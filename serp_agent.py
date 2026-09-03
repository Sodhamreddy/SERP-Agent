#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
AHNS SERP Position Tracker — myassuredhomenursing.com

Architecture : AGENTIC (locked) — Ollama coordinator + Search Provider.
               Flow: Ollama plan → Gemini Google-Search grounding →
               SERP parser → ranking match → Excel/CSV report.
               No Playwright / browser automation.
Position     : PAGE.POSITION  (e.g. 2.3 = page 2, result 3) — approximate,
               based on the order of sources the search provider returns.

Setup (one time):
    pip install -r requirements.txt

Run:
    python serp_agent.py          (CLI)
    double-click start_ui.bat     (Web UI)
"""
from __future__ import annotations

import os
import re
import sys
import json
import logging
from datetime import datetime


import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side

# ───────────────────────────────────────────────────────────────
# CONFIGURATION
# ───────────────────────────────────────────────────────────────
_DEFAULT_DOMAIN = "myassuredhomenursing.com"   # seed; overridden by config.json
COMPETITOR_LIMIT = 5      # competitor SERP results saved per keyword for AI analysis
PAGES_TO_CHECK  = 5       # fast test mode: 5 pages × 10 results = top 50 positions
KEYWORD_LIMIT   = None    # None = all 120 keywords; set a number for test runs
# Batch mode: max keywords to attempt per run. Free search tiers (Ollama session,
# Gemini daily) can't do all 120 at once, so each run does up to BATCH_SIZE of the
# remaining keywords and checkpoints progress. Re-run after the limit resets to
# continue; the final report is written automatically once all keywords are done.
# Set to 0 (or None) to disable batching and attempt every keyword in one run.
BATCH_SIZE      = int(os.getenv("SERP_BATCH_SIZE", "40"))
PROGRESS_PATH   = os.path.join("results", "_progress.json")


logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s  %(message)s",
    datefmt="%H:%M:%S",
    handlers=[
        logging.StreamHandler(sys.stdout),
        logging.FileHandler("serp_agent.log", encoding="utf-8"),
    ],
)
log = logging.getLogger(__name__)


# ───────────────────────────────────────────────────────────────
# KEYWORDS  — default seed (30 cities × 4 keyword types). The live list is
# loaded from keywords.json and is editable at runtime via the Settings UI.
# ───────────────────────────────────────────────────────────────
_DEFAULT_KEYWORDS: list[tuple[str, str]] = [
    # Birmingham, MI
    ("Birmingham, MI",              "Senior In home care in Birmingham,Mi"),
    ("Birmingham, MI",              "home care services in Birmingham,Mi"),
    ("Birmingham, MI",              "In home care services in Birmingham,Mi"),
    ("Birmingham, MI",              "24 hour home care in Birmingham,Mi"),
    # Beverly Hills, MI
    ("Beverly Hills, MI",           "Senior In home care in Beverly Hills,Mi"),
    ("Beverly Hills, MI",           "home care services in Beverly Hills,Mi"),
    ("Beverly Hills, MI",           "In home care services in Beverly Hills,Mi"),
    ("Beverly Hills, MI",           "24 hour home care in Beverly Hills,Mi"),
    # Ann Arbor, MI
    ("Ann Arbor, MI",               "Senior In home care in Ann Arbor,Mi"),
    ("Ann Arbor, MI",               "home care services in Ann Arbor,Mi"),
    ("Ann Arbor, MI",               "In home care services in Ann Arbor,Mi"),
    ("Ann Arbor, MI",               "24 hour home care in Ann Arbor,Mi"),
    # Bloomfield Hills, MI
    ("Bloomfield Hills, MI",        "Senior In home care in Bloomfield Hills,Mi"),
    ("Bloomfield Hills, MI",        "home care services in Bloomfield Hills,Mi"),
    ("Bloomfield Hills, MI",        "In home care services in Bloomfield Hills,Mi"),
    ("Bloomfield Hills, MI",        "24 hour home care in Bloomfield Hills,Mi"),
    # Clinton Township, MI
    ("Clinton Township, MI",        "Senior In home care in Clinton Township,MI"),
    ("Clinton Township, MI",        "home care services in Clinton Township,MI"),
    ("Clinton Township, MI",        "In home care services in Clinton Township,MI"),
    ("Clinton Township, MI",        "24 hour home care in Clinton Township,MI"),
    # Dearborn, MI
    ("Dearborn, MI",                "Senior In home care in Dearborn, MI"),
    ("Dearborn, MI",                "home care services in Dearborn, MI"),
    ("Dearborn, MI",                "In home care services in Dearborn, MI"),
    ("Dearborn, MI",                "24 hour home care in Dearborn, MI"),
    # Detroit, MI
    ("Detroit, MI",                 "Senior In home care in Detroit,MI"),
    ("Detroit, MI",                 "home care services in Detroit,MI"),
    ("Detroit, MI",                 "In home care services in Detroit,MI"),
    ("Detroit, MI",                 "24 hour home care in Detroit,MI"),
    # Canton, MI
    ("Canton, MI",                  "Senior In home care in Canton,MI"),
    ("Canton, MI",                  "home care services in Canton,MI"),
    ("Canton, MI",                  "In home care services in Canton,MI"),
    ("Canton, MI",                  "24 hour home care in Canton,MI"),
    # Franklin, MI
    ("Franklin, MI",                "Senior In home care in Franklin,MI"),
    ("Franklin, MI",                "home care services in Franklin,MI"),
    ("Franklin, MI",                "In home care services in Franklin,MI"),
    ("Franklin, MI",                "24 hour home care in Franklin,MI"),
    # Fraser, MI
    ("Fraser, MI",                  "Senior In home care in Fraser,MI"),
    ("Fraser, MI",                  "home care services in Fraser,MI"),
    ("Fraser, MI",                  "In home care services in Fraser,MI"),
    ("Fraser, MI",                  "24 hour home care in Fraser,MI"),
    # Grosse Pointe, MI
    ("Grosse Pointe, MI",           "Senior In home care in Grosse Pointe, MI"),
    ("Grosse Pointe, MI",           "home care services in Grosse Pointe, MI"),
    ("Grosse Pointe, MI",           "In home care services in Grosse Pointe, MI"),
    ("Grosse Pointe, MI",           "24 hour home care in Grosse Pointe, MI"),
    # Novi, MI
    ("Novi, MI",                    "Senior In home care in Novi,Mi"),
    ("Novi, MI",                    "home care services in Novi,Mi"),
    ("Novi, MI",                    "In home care services in Novi,Mi"),
    ("Novi, MI",                    "24 hour home care in Novi,Mi"),
    # Northville, MI
    ("Northville, MI",              "Senior In home care in Northville,Mi"),
    ("Northville, MI",              "home care services in Northville,Mi"),
    ("Northville, MI",              "In home care services in Northville,Mi"),
    ("Northville, MI",              "24 hour home care in Northville,Mi"),
    # Rochester Hills, MI
    ("Rochester Hills, MI",         "Senior In home care in Rochester Hills, MI"),
    ("Rochester Hills, MI",         "home care services in Rochester Hills, MI"),
    ("Rochester Hills, MI",         "In home care services in Rochester Hills, MI"),
    ("Rochester Hills, MI",         "24 hour home care in Rochester Hills, MI"),
    # Royal Oak, MI
    ("Royal Oak, MI",               "Senior In home care in Royal Oak,MI"),
    ("Royal Oak, MI",               "home care services in Royal Oak,MI"),
    ("Royal Oak, MI",               "In home care services in Royal Oak,MI"),
    ("Royal Oak, MI",               "24 hour home care in Royal Oak,MI"),
    # Southfield, MI
    ("Southfield, MI",              "Senior In home care in Southfield, MI"),
    ("Southfield, MI",              "home care services in Southfield, MI"),
    ("Southfield, MI",              "In home care services in Southfield, MI"),
    ("Southfield, MI",              "24 hour home care in Southfield, MI"),
    # Troy, MI
    ("Troy, MI",                    "Senior In home care in Troy,MI"),
    ("Troy, MI",                    "home care services in Troy,MI"),
    ("Troy, MI",                    "In home care services in Troy,MI"),
    ("Troy, MI",                    "24 hour home care in Troy,MI"),
    # Madison Heights, MI
    ("Madison Heights, MI",         "Senior In home care in Madison Heights,MI"),
    ("Madison Heights, MI",         "home care services in Madison Heights,MI"),
    ("Madison Heights, MI",         "In home care services in Madison Heights,MI"),
    ("Madison Heights, MI",         "24 hour home care in Madison Heights,MI"),
    # Bingham Farms, MI
    ("Bingham Farms, MI",           "Senior In home care in Bingham Farms,MI"),
    ("Bingham Farms, MI",           "home care services in Bingham Farms,MI"),
    ("Bingham Farms, MI",           "In home care services in Bingham Farms,MI"),
    ("Bingham Farms, MI",           "24 hour home care in Bingham Farms,MI"),
    # Macomb Township, MI
    ("Macomb Township, MI",         "Senior In home care in Macomb Township,MI"),
    ("Macomb Township, MI",         "home care services in Macomb Township,MI"),
    ("Macomb Township, MI",         "In home care services in Macomb Township,MI"),
    ("Macomb Township, MI",         "24 hour home care in Macomb Township,MI"),
    # Waterford, MI
    ("Waterford, MI",               "Senior In home care in Waterford,MI"),
    ("Waterford, MI",               "home care services in Waterford,MI"),
    ("Waterford, MI",               "In home care services in Waterford,MI"),
    ("Waterford, MI",               "24 hour home care in Waterford,MI"),
    # Farmington Hills, MI
    ("Farmington Hills, MI",        "Senior In home care in Farmington Hills,MI"),
    ("Farmington Hills, MI",        "home care services in Farmington Hills,MI"),
    ("Farmington Hills, MI",        "In home care services in Farmington Hills, MI"),
    ("Farmington Hills, MI",        "24 hour home care in Farmington Hills,MI"),
    # Shelby Township, MI
    ("Shelby Township, MI",         "Senior In home care in Shelby Township,MI"),
    ("Shelby Township, MI",         "home care services in Shelby Township,MI"),
    ("Shelby Township, MI",         "In home care services in Shelby Township,MI"),
    ("Shelby Township, MI",         "24 hour home care in Shelby Township,MI"),
    # St Clair Shores, MI
    ("St Clair Shores, MI",         "Senior In home care in St Clair Shores,MI"),
    ("St Clair Shores, MI",         "home care services in St Clair Shores,MI"),
    ("St Clair Shores, MI",         "In home care services in St Clair Shores,MI"),
    ("St Clair Shores, MI",         "24 hour home care in St Clair Shores,MI"),
    # Sterling Heights, MI
    ("Sterling Heights, MI",        "Senior In home care in Sterling Heights,MI"),
    ("Sterling Heights, MI",        "home care services in Sterling Heights,MI"),
    ("Sterling Heights, MI",        "In home care services in Sterling Heights,MI"),
    ("Sterling Heights, MI",        "24 hour home care in Sterling Heights,MI"),
    # West Bloomfield, MI
    ("West Bloomfield, MI",         "Senior In home care in West Bloomfield,Mi"),
    ("West Bloomfield, MI",         "home care services in West Bloomfield, MI"),
    ("West Bloomfield, MI",         "In home care services in West Bloomfield,Mi"),
    ("West Bloomfield, MI",         "24 hour home care in West Bloomfield,Mi"),
    # Warren, MI
    ("Warren, MI",                  "Senior In home care in Warren, Mi"),
    ("Warren, MI",                  "home care services in Warren, Mi"),
    ("Warren, MI",                  "In home care services in Warren, Mi"),
    ("Warren, MI",                  "24 hour home care in Warren, Mi"),
    # Milford Charter Township, MI
    ("Milford Charter Township, MI","Senior In home care in Milford Charter Township, MI"),
    ("Milford Charter Township, MI","home care services in Milford Charter Township, MI"),
    ("Milford Charter Township, MI","In home care services in Milford Charter Township, MI"),
    ("Milford Charter Township, MI","24 hour home care in Milford Charter Township, MI"),
    # South Lyon, MI
    ("South Lyon, MI",              "Senior In home care in South Lyon, Mi"),
    ("South Lyon, MI",              "home care services in South Lyon, Mi"),
    ("South Lyon, MI",              "In home care services in South Lyon, Mi"),
    ("South Lyon, MI",              "24 hour home care in South Lyon, Mi"),
    # Plymouth, MI
    ("Plymouth, MI",                "Senior In home care in Plymouth, Mi"),
    ("Plymouth, MI",                "home care services in Plymouth, Mi"),
    ("Plymouth, MI",                "In home care services in Plymouth, Mi"),
    ("Plymouth, MI",                "24 hour home care in Plymouth, Mi"),
]


# ───────────────────────────────────────────────────────────────
# CLIENT PROFILES — multiple saved clients (each: name, domain, keywords),
# stored in clients.json, switchable at runtime via the Settings UI. The active
# client's domain + keywords are exposed as live module attributes
# (TARGET_DOMAIN / KEYWORDS) so updates propagate everywhere with no restart.
# ───────────────────────────────────────────────────────────────
_CONFIG_PATH    = os.path.join(os.path.dirname(os.path.abspath(__file__)), "config.json")
KEYWORDS_PATH   = os.path.join(os.path.dirname(os.path.abspath(__file__)), "keywords.json")
CLIENTS_PATH    = os.path.join(os.path.dirname(os.path.abspath(__file__)), "clients.json")


def _read_config() -> dict:
    try:
        with open(_CONFIG_PATH, "r", encoding="utf-8") as fh:
            data = json.load(fh)
        return data if isinstance(data, dict) else {}
    except Exception:
        return {}


def _slugify(name: str) -> str:
    slug = "".join(c.lower() if c.isalnum() else "-" for c in name).strip("-")
    while "--" in slug:
        slug = slug.replace("--", "-")
    return slug or "client"


def _normalize_pairs(pairs) -> list[list[str]]:
    out: list[list[str]] = []
    for item in pairs or []:
        if isinstance(item, (list, tuple)) and len(item) >= 2:
            loc, kw = str(item[0]).strip(), str(item[1]).strip()
        elif isinstance(item, str):
            loc, kw = "", item.strip()
        else:
            continue
        if kw:
            out.append([loc, kw])
    return out


def _seed_clients() -> dict:
    """Build the initial clients store, migrating any existing single-client config.

    The seed AHNS client starts ownerless (owner="") — the first admin to sign up
    claims all ownerless clients (see assign_ownerless_to).
    """
    domain = str(_read_config().get("TARGET_DOMAIN", "")).strip() or _DEFAULT_DOMAIN
    pairs: list = []
    try:
        with open(KEYWORDS_PATH, "r", encoding="utf-8") as fh:
            pairs = _normalize_pairs(json.load(fh))
    except Exception:
        pairs = []
    if not pairs:
        pairs = _normalize_pairs(_DEFAULT_KEYWORDS)
    return {
        "active": "ahns",
        "clients": [
            {"id": "ahns", "name": "Assured Home Nursing", "domain": domain,
             "keywords": pairs, "owner": "", "created_at": ""}
        ],
    }


def load_clients() -> dict:
    """Load clients.json; migrate/seed on first run. Always returns a valid store."""
    try:
        with open(CLIENTS_PATH, "r", encoding="utf-8") as fh:
            data = json.load(fh)
        if isinstance(data, dict) and data.get("clients"):
            for c in data["clients"]:        # ensure ownership fields exist
                c.setdefault("owner", "")
                c.setdefault("created_at", "")
            return data
    except Exception:
        pass
    data = _seed_clients()
    save_clients(data)
    return data


def save_clients(data: dict) -> dict:
    with open(CLIENTS_PATH, "w", encoding="utf-8") as fh:
        json.dump(data, fh, ensure_ascii=False, indent=2)
    return data


def get_client(client_id: str) -> dict | None:
    for c in load_clients()["clients"]:
        if c["id"] == client_id:
            return c
    return None


def _is_admin(user: dict | None) -> bool:
    return bool(user) and user.get("role") == "admin"


def can_access(user: dict | None, client: dict | None) -> bool:
    """Admins access everything (incl. ownerless); members only their own."""
    if not client:
        return False
    if _is_admin(user):
        return True
    return bool(user) and client.get("owner") == user.get("email")


def clients_for(user: dict | None) -> list[dict]:
    """All clients for an admin; only the member's own for a member."""
    everything = load_clients()["clients"]
    if _is_admin(user):
        return everything
    email = (user or {}).get("email")
    return [c for c in everything if c.get("owner") == email]


def assign_ownerless_to(email: str) -> None:
    """Claim all ownerless clients for `email` (used when the first admin signs up)."""
    data = load_clients()
    changed = False
    for c in data["clients"]:
        if not c.get("owner"):
            c["owner"] = email
            changed = True
    if changed:
        save_clients(data)


def default_keywords(client_id: str | None = None) -> list[tuple[str, str]]:
    c = get_client(client_id) if client_id else None
    c = c or load_clients()["clients"][0]
    return [(loc, kw) for loc, kw in _normalize_pairs(c.get("keywords"))]


def upsert_client(owner: str, name: str, domain: str, keywords,
                  client_id: str | None = None) -> dict:
    """Create or update a client (owner-stamped). Returns the saved client dict."""
    import websearch_agent
    name = (name or "").strip() or "Client"
    clean_domain = websearch_agent._normalize_domain(domain) or (domain or "").strip()
    if not clean_domain:
        raise ValueError("Domain cannot be empty.")
    pairs = _normalize_pairs(keywords)

    data = load_clients()
    if not client_id:
        base = _slugify(name)
        client_id, n = base, 2
        existing = {c["id"] for c in data["clients"]}
        while client_id in existing:
            client_id, n = f"{base}-{n}", n + 1

    found = next((c for c in data["clients"] if c["id"] == client_id), None)
    if found:
        found.update({"name": name, "domain": clean_domain, "keywords": pairs})
        if not found.get("owner"):
            found["owner"] = owner
    else:
        data["clients"].append({
            "id": client_id, "name": name, "domain": clean_domain, "keywords": pairs,
            "owner": owner, "created_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        })
    save_clients(data)
    return next(c for c in data["clients"] if c["id"] == client_id)


def delete_client(client_id: str) -> None:
    data = load_clients()
    data["clients"] = [c for c in data["clients"] if c["id"] != client_id]
    if data.get("active") == client_id:
        data["active"] = data["clients"][0]["id"] if data["clients"] else ""
    save_clients(data)


# ── CLI default (web paths resolve the active client per-request from the session) ──
def _default_client() -> dict:
    data = load_clients()
    return next((c for c in data["clients"] if c["id"] == data.get("active")),
                data["clients"][0])


_dc = _default_client()
TARGET_DOMAIN = _dc.get("domain") or _DEFAULT_DOMAIN
KEYWORDS: list[tuple[str, str]] = [(loc, kw) for loc, kw in _normalize_pairs(_dc.get("keywords"))]
CLIENT_NAME = _dc.get("name") or "SERP Agent"


def report_label(name: str | None) -> str:
    """Filename-safe client label — the "<Client> SERP <ts>" report prefix.

    Keeping the client in the filename means the sidebar AND the downloaded
    file both say who the report is for.
    """
    safe = re.sub(r"[^A-Za-z0-9 &._-]", " ", str(name or ""))
    safe = re.sub(r"\s+", " ", safe).strip()
    safe = " ".join(w for w in safe.split() if w.upper() != "SERP")   # keep the split unambiguous
    return safe[:50].strip() or "SERP Agent"


def run_serp_agent(keywords: list[tuple[str, str]], domain: str | None = None) -> list[dict]:
    """
    Agentic SERP entry point — LOCKED to agentic mode.

    Flow (always):
        Coordinator -> Search Provider -> SERP Parser -> Report

    `domain` is the target site to match (defaults to the module's default client).
    """
    from websearch_agent import run_websearch_agent

    target = domain or TARGET_DOMAIN
    log.info("SERP agent mode: AGENTIC — target %s", target)
    return run_websearch_agent(keywords, target, log)


# ───────────────────────────────────────────────────────────────
# EXCEL OUTPUT  (Summary sheet + grouped SERP Report sheet)
# ───────────────────────────────────────────────────────────────

# Professional palette: slate headers, a single clean accent blue, and SOFT status
# tints (light green / light red) so the dark text stays readable instead of being
# flooded by saturated fills. Status is reinforced with colored text in the data rows.
_BLUE_FILL   = PatternFill("solid", fgColor="2563EB")   # accent blue (section headers)
_GREEN_FILL  = PatternFill("solid", fgColor="DCFCE7")   # soft green tint (positive)
_RED_FILL    = PatternFill("solid", fgColor="FEE2E2")   # soft red tint (negative)
_GREY_FILL   = PatternFill("solid", fgColor="F1F5F9")   # slate-100 (sub-headers)
_SUM_HEAD    = PatternFill("solid", fgColor="0F172A")   # slate-900 (titles / totals)
_WHITE_FONT  = Font(bold=True, color="FFFFFF", size=11)
_BOLD_FONT   = Font(bold=True, size=10, color="1E293B")
_NORMAL_FONT = Font(size=10, color="1E293B")
_GREEN_FONT  = Font(size=10, bold=True, color="15803D")  # "Yes" reviews
_RED_FONT    = Font(size=10, bold=True, color="B91C1C")  # "No" reviews
_CENTER      = Alignment(horizontal="center", vertical="center")
_LEFT        = Alignment(horizontal="left",   vertical="center")
_THIN        = Side(border_style="thin", color="E2E8F0")
_BORDER      = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)


def _cell(ws, row, col, value="", fill=None, font=None, align=None):
    c = ws.cell(row=row, column=col, value=value)
    if fill:  c.fill = fill
    if font:  c.font = font
    if align: c.alignment = align
    c.border = _BORDER
    return c


def save_excel(df: pd.DataFrame, path: str, run_date: str, domain: str | None = None,
               comparison: dict | None = None) -> None:
    wb = Workbook()

    # ── Sheet 1: Executive Summary ────────────────────────────
    ws_s = wb.active
    ws_s.title = "Summary"
    ws_s.column_dimensions["A"].width = 30
    ws_s.column_dimensions["B"].width = 18
    ws_s.column_dimensions["C"].width = 18
    ws_s.column_dimensions["D"].width = 18
    ws_s.column_dimensions["E"].width = 18

    # Title row
    ws_s.merge_cells("A1:E1")
    t = ws_s.cell(row=1, column=1, value=f"SERP Report — {domain or TARGET_DOMAIN}")
    t.fill = _SUM_HEAD
    t.font = Font(bold=True, color="FFFFFF", size=14)
    t.alignment = _CENTER
    ws_s.row_dimensions[1].height = 30

    # Run info
    ws_s.merge_cells("A2:E2")
    d = ws_s.cell(row=2, column=1, value=f"Run Date: {run_date}   |   Total Keywords: {len(df)}")
    d.fill = PatternFill("solid", fgColor="1E293B")
    d.font = Font(color="FFFFFF", size=10)
    d.alignment = _CENTER
    ws_s.row_dimensions[2].height = 18

    # Overall stats header
    row = 4
    for col, label in enumerate(["Location", "Total", "Ranked", "Page 1", "Has Reviews"], 1):
        _cell(ws_s, row, col, label, fill=_BLUE_FILL, font=_WHITE_FONT, align=_CENTER)
    ws_s.row_dimensions[row].height = 18
    row += 1

    total_all, ranked_all, p1_all, rev_all = 0, 0, 0, 0
    for location, grp in df.groupby("Location", sort=False):
        total   = len(grp)
        ranked  = (grp["Position"] != "Not Found").sum()
        page1   = grp["Position"].apply(
            lambda x: str(x).startswith("1.") if x != "Not Found" else False
        ).sum()
        reviews = (grp["Reviews"] == "Yes").sum()
        total_all += total; ranked_all += ranked
        p1_all += page1;    rev_all += reviews

        fill = _GREEN_FILL if ranked > 0 else _RED_FILL
        for col, val in enumerate([location, total, ranked, page1, reviews], 1):
            _cell(ws_s, row, col, val, fill=fill, font=_NORMAL_FONT,
                  align=_LEFT if col == 1 else _CENTER)
        row += 1

    # Totals row
    _cell(ws_s, row, 1, "TOTAL", fill=_SUM_HEAD, font=_WHITE_FONT, align=_CENTER)
    for col, val in enumerate([total_all, ranked_all, p1_all, rev_all], 2):
        _cell(ws_s, row, col, val, fill=_SUM_HEAD, font=_WHITE_FONT, align=_CENTER)

    # ── Sheet 2: Full SERP Report ──────────────────────────────
    ws = wb.create_sheet("SERP Report")
    ws.column_dimensions["A"].width = 6
    ws.column_dimensions["B"].width = 52
    ws.column_dimensions["C"].width = 14
    ws.column_dimensions["D"].width = 12
    ws.column_dimensions["E"].width = 22

    row_idx = 1
    for location, grp in df.groupby("Location", sort=False):
        # Location header (merged A:B)
        ws.merge_cells(
            start_row=row_idx, start_column=1,
            end_row=row_idx,   end_column=2
        )
        h = ws.cell(row=row_idx, column=1, value=f"Keywords — {location}")
        h.fill = _BLUE_FILL; h.font = _WHITE_FONT
        h.alignment = _LEFT; h.border = _BORDER

        d_cell = ws.cell(row=row_idx, column=3, value=run_date)
        d_cell.fill = _BLUE_FILL; d_cell.font = _WHITE_FONT
        d_cell.alignment = _CENTER; d_cell.border = _BORDER
        for col in [4, 5]:
            c = ws.cell(row=row_idx, column=col, value="")
            c.fill = _BLUE_FILL; c.border = _BORDER
        row_idx += 1

        # Sub-header
        for col, label in enumerate(["S.No", "Keyword", "Position", "Reviews", "Checked At"], 1):
            _cell(ws, row_idx, col, label, fill=_GREY_FILL, font=_BOLD_FONT, align=_CENTER)
        ws.row_dimensions[row_idx].height = 16
        row_idx += 1

        # Data rows
        for sno, (_, r) in enumerate(grp.iterrows(), 1):
            has_rev = r["Reviews"] == "Yes"
            fill = _GREEN_FILL if has_rev else _RED_FILL
            vals = [sno, r["Keyword"], r["Position"], r["Reviews"], r["Checked At"]]
            for col, val in enumerate(vals, 1):
                _cell(ws, row_idx, col, val, fill=fill,
                      font=_NORMAL_FONT, align=_LEFT)
            ws.cell(row=row_idx, column=1).alignment = _CENTER
            ws.cell(row=row_idx, column=3).alignment = _CENTER
            # Reviews cell: colored, bold text reinforces the status on the soft tint.
            rev_cell = ws.cell(row=row_idx, column=4)
            rev_cell.alignment = _CENTER
            rev_cell.font = _GREEN_FONT if has_rev else _RED_FONT
            ws.row_dimensions[row_idx].height = 15
            row_idx += 1

        row_idx += 1  # blank separator

    # Auto-filter on SERP Report sheet
    ws.auto_filter.ref = f"A2:E{row_idx - 1}"

    # ── Sheet 3: vs Last Run (per-city comparison with the previous run) ──
    if comparison and comparison.get("locations"):
        ws_c = wb.create_sheet("vs Last Run")
        for col, w in {"A": 26, "B": 50, "C": 13, "D": 13, "E": 12, "F": 10, "G": 8, "H": 8}.items():
            ws_c.column_dimensions[col].width = w

        ws_c.merge_cells("A1:H1")
        t = ws_c.cell(row=1, column=1,
                      value=f"This run vs previous run ({comparison.get('previous_run', '')})")
        t.fill = _SUM_HEAD
        t.font = Font(bold=True, color="FFFFFF", size=12)
        t.alignment = _CENTER
        ws_c.row_dimensions[1].height = 24

        # Per-city summary table
        row = 3
        headers = ["Location", "Keywords", "Ranked (last)", "Ranked (now)",
                   "Improved", "Declined", "New", "Lost"]
        for col, label in enumerate(headers, 1):
            _cell(ws_c, row, col, label, fill=_BLUE_FILL, font=_WHITE_FONT, align=_CENTER)
        row += 1
        for e in comparison["locations"]:
            up   = e.get("improved", 0) + e.get("new_ranked", 0)
            down = e.get("declined", 0) + e.get("lost", 0)
            fill = _GREEN_FILL if up > down else (_RED_FILL if down > up else _GREY_FILL)
            vals = [e.get("location", ""), e.get("checked", 0), e.get("ranked_prev", 0),
                    e.get("ranked_now", 0), e.get("improved", 0), e.get("declined", 0),
                    e.get("new_ranked", 0), e.get("lost", 0)]
            for col, val in enumerate(vals, 1):
                _cell(ws_c, row, col, val, fill=fill, font=_NORMAL_FONT,
                      align=_LEFT if col == 1 else _CENTER)
            row += 1

        # Keyword-level movements
        moves = [(e.get("location", ""), c)
                 for e in comparison["locations"] for c in e.get("changes", [])]
        if moves:
            row += 1
            ws_c.merge_cells(start_row=row, start_column=1, end_row=row, end_column=8)
            h = ws_c.cell(row=row, column=1, value="Keyword movements")
            h.fill = _SUM_HEAD; h.font = _WHITE_FONT
            h.alignment = _LEFT; h.border = _BORDER
            row += 1
            for col, label in enumerate(["Location", "Keyword", "Last", "Now", "Change"], 1):
                _cell(ws_c, row, col, label, fill=_GREY_FILL, font=_BOLD_FONT, align=_CENTER)
            row += 1
            for loc, c in moves:
                worse = c.get("delta") == "lost" or str(c.get("delta", "")).startswith("-")
                vals = [loc, c.get("keyword", ""), c.get("last", ""), c.get("now", ""),
                        c.get("delta", "")]
                for col, val in enumerate(vals, 1):
                    _cell(ws_c, row, col, val, fill=_RED_FILL if worse else _GREEN_FILL,
                          font=_NORMAL_FONT, align=_LEFT if col <= 2 else _CENTER)
                d_cell = ws_c.cell(row=row, column=5)
                d_cell.font = _RED_FONT if worse else _GREEN_FONT
                row += 1

    wb.save(path)
    log.info("Excel saved: %s", path)


# ───────────────────────────────────────────────────────────────
# BATCH PROGRESS (checkpoint so runs can resume across reset windows)
# ───────────────────────────────────────────────────────────────

def _kw_key(location: str, keyword: str) -> str:
    return f"{location}||{keyword}"


def _load_progress() -> dict[str, dict]:
    """Return completed rows keyed by location||keyword (empty if none)."""
    try:
        with open(PROGRESS_PATH, encoding="utf-8") as fh:
            data = json.load(fh)
        if isinstance(data, list):
            return {_kw_key(r.get("Location", ""), r.get("Keyword", "")): r for r in data}
    except Exception:
        pass
    return {}


def _save_progress(done: dict[str, dict]) -> None:
    os.makedirs("results", exist_ok=True)
    with open(PROGRESS_PATH, "w", encoding="utf-8") as fh:
        json.dump(list(done.values()), fh, ensure_ascii=False, indent=2)


def _write_report(rows: list[dict], run_date: str, ts: str) -> tuple[str, str]:
    df = pd.DataFrame([{k: v for k, v in r.items() if k != "_ok"} for r in rows])
    os.makedirs("results", exist_ok=True)
    label      = report_label(CLIENT_NAME)
    excel_path = f"results/{label} SERP {ts}.xlsx"
    csv_path   = f"results/{label} SERP {ts}.csv"
    save_excel(df, excel_path, run_date)
    df.to_csv(csv_path, index=False)

    ranked  = (df["Position"] != "Not Found").sum()
    reviews = (df["Reviews"] == "Yes").sum()
    total   = len(df)
    print(f"\n{'=' * 65}")
    print(f"  FINAL REPORT  {datetime.now():%d-%m-%Y %H:%M:%S}")
    print(f"  Ranked   : {ranked}/{total}  ({ranked / total * 100:.1f}%)")
    print(f"  Reviews  : {reviews}/{total}")
    print(f"  Excel    : {excel_path}")
    print(f"  CSV      : {csv_path}")
    print(f"{'=' * 65}\n")
    return excel_path, csv_path


# ───────────────────────────────────────────────────────────────
# MAIN
# ───────────────────────────────────────────────────────────────

def main() -> None:
    run_date = datetime.now().strftime("%d/%m/%Y")
    ts       = datetime.now().strftime("%d-%m-%Y %H-%M-%S")

    all_keywords = KEYWORDS[:KEYWORD_LIMIT] if KEYWORD_LIMIT else KEYWORDS
    total = len(all_keywords)

    done = _load_progress()
    # Keep only progress entries that belong to the current keyword set.
    done = {k: v for k, v in done.items()
            if k in {_kw_key(loc, kw) for loc, kw in all_keywords}}
    remaining = [(loc, kw) for loc, kw in all_keywords
                 if _kw_key(loc, kw) not in done]

    print(f"\n{'=' * 65}")
    print(f"  AHNS SERP Tracker — {TARGET_DOMAIN}")
    print(f"  Engine   : Ollama + Search Provider")
    print(f"  Progress : {len(done)}/{total} done, {len(remaining)} remaining")
    print(f"  Started  : {datetime.now():%d-%m-%Y %H:%M:%S}")
    print(f"{'=' * 65}\n")

    # Everything already complete → emit the final report and archive progress.
    if not remaining:
        rows = [done[_kw_key(loc, kw)] for loc, kw in all_keywords]
        _write_report(rows, run_date, ts)
        try:
            os.replace(PROGRESS_PATH, PROGRESS_PATH.replace(".json", f"_{ts}.done.json"))
        except Exception:
            pass
        print("  All keywords complete. Progress archived.\n")
        return

    # Run up to BATCH_SIZE of the remaining keywords this session.
    batch = remaining[:BATCH_SIZE] if BATCH_SIZE else remaining
    log.info("Running batch of %d keyword(s) (%d still remaining after the batch cap)",
             len(batch), max(0, len(remaining) - len(batch)))

    rows = run_serp_agent(batch)

    ok_rows   = [r for r in rows if r.get("_ok")]
    failed    = [r for r in rows if not r.get("_ok")]
    for r in ok_rows:
        done[_kw_key(r["Location"], r["Keyword"])] = r
    _save_progress(done)

    still_remaining = total - len(done)
    print(f"\n{'=' * 65}")
    print(f"  BATCH DONE  {datetime.now():%d-%m-%Y %H:%M:%S}")
    print(f"  This batch : {len(ok_rows)} succeeded, {len(failed)} rate-limited/failed")
    print(f"  Overall    : {len(done)}/{total} complete, {still_remaining} remaining")
    print(f"  Progress   : {PROGRESS_PATH}")
    print(f"{'=' * 65}")

    if still_remaining > 0:
        print("  NOT finished. Re-run `python serp_agent.py` after the search")
        print("  provider's limit resets (Ollama session ~3h) to continue.\n")
    else:
        # The batch completed the set — write the final report now.
        ordered = [done[_kw_key(loc, kw)] for loc, kw in all_keywords]
        _write_report(ordered, run_date, ts)
        try:
            os.replace(PROGRESS_PATH, PROGRESS_PATH.replace(".json", f"_{ts}.done.json"))
        except Exception:
            pass
        print("  All keywords complete. Progress archived.\n")


if __name__ == "__main__":
    main()
